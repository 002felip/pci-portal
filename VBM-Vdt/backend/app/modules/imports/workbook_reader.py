from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Optional

import openpyxl

from app.modules.imports.issues import IssueCollector
from app.modules.imports.models import (
    ParsedNode,
    ParsedTree,
    ParsedValue,
    ParsedWorkbook,
    YearBlock,
)

METADATA_SHEET = "01_Workbook"
REGISTRY_SHEET = "02_Trees"
REGISTRY_FIRST_DATA_ROW = 5
NODE_FIRST_DATA_ROW = 3
STRUCTURAL_COLUMN_COUNT = 9
CONTEXTS = ["Actuals", "Budget", "LoBP", "Benchmark"]

_METADATA_ROWS = {
    5: "template_version",
    6: "workbook_type",
    7: "region_code",
    8: "region_name",
    9: "revision_cycle_year",
}


def _text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _int(value) -> Optional[int]:
    text = _text(value)
    if text is None:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _decimal(value, sheet, row, column, issues: IssueCollector):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        issues.blocking(
            "VALUE_NOT_NUMERIC",
            f"Value {value!r} is not a number",
            sheet=sheet, row=row, column=str(column),
        )
        return None


def read_workbook(path, issues: IssueCollector) -> ParsedWorkbook:
    # read_only keeps the file handle open until close(); on Windows that
    # blocks deleting the upload's temp copy, and it leaks handles anywhere.
    wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=False)
    try:
        return _read(wb, issues)
    finally:
        wb.close()


def _read(wb, issues: IssueCollector) -> ParsedWorkbook:
    parsed = ParsedWorkbook()

    if METADATA_SHEET not in wb.sheetnames:
        issues.blocking("MISSING_SHEET", f"Required sheet {METADATA_SHEET!r} is missing",
                        sheet=METADATA_SHEET)
    else:
        ws = wb[METADATA_SHEET]
        rows = {r: list(row) for r, row in enumerate(
            ws.iter_rows(min_row=1, max_row=12, min_col=1, max_col=3), start=1)}
        for row, attr in _METADATA_ROWS.items():
            cells = rows.get(row, [None, None, None])
            cell = cells[2] if len(cells) > 2 else None
            raw = cell.value if cell is not None else None
            setattr(parsed, attr, _int(raw) if attr == "revision_cycle_year" else _text(raw))

    if REGISTRY_SHEET not in wb.sheetnames:
        issues.blocking("MISSING_SHEET", f"Required sheet {REGISTRY_SHEET!r} is missing",
                        sheet=REGISTRY_SHEET)
        return parsed

    ws = wb[REGISTRY_SHEET]
    for r, row in enumerate(ws.iter_rows(min_row=REGISTRY_FIRST_DATA_ROW, min_col=1, max_col=5),
                            start=REGISTRY_FIRST_DATA_ROW):
        values = [_text(c.value) for c in row]
        if not any(values):
            continue
        tree = ParsedTree(
            tree_code=values[0], tree_name=values[1], tree_type=values[2],
            pair_code=values[3], worksheet=values[4], registry_row=r,
        )
        parsed.trees.append(tree)
        if tree.worksheet and tree.worksheet not in wb.sheetnames:
            issues.blocking(
                "WORKSHEET_NOT_FOUND",
                f"Tree {tree.tree_code!r} points at worksheet {tree.worksheet!r}, which does not exist",
                sheet=REGISTRY_SHEET, row=r, column="E",
            )
            continue
        if tree.worksheet:
            _read_tree_sheet(wb[tree.worksheet], tree, issues)

    return parsed


def _read_year_blocks(rows, tree: ParsedTree, issues: IssueCollector) -> None:
    # A year's columns don't all need to be present, and don't need a fixed
    # order - a tree can import e.g. only "LoBP 2027" and "Budget 2027" and
    # skip Actuals/Benchmark for that year entirely, or import a lone
    # "LoBP 2030" with no other context for that year at all. Every
    # populated column is validated independently and grouped into its
    # year's block by context.
    year_row = rows.get(1, {})
    context_row = rows.get(2, {})
    columns = sorted(c for c, v in year_row.items()
                     if c >= STRUCTURAL_COLUMN_COUNT + 1 and v is not None)
    if not columns:
        issues.blocking("NO_YEAR_BLOCKS", "No annual value columns found",
                        sheet=tree.worksheet, row=1)
        return

    blocks_by_year: dict[int, YearBlock] = {}
    for c in columns:
        year = _int(year_row.get(c))
        if year is None or not (1900 <= year <= 2999):
            issues.blocking(
                "INVALID_YEAR_HEADER",
                f"Row 1 column {c} must hold a four-digit year, found {year_row.get(c)!r}",
                sheet=tree.worksheet, row=1, column=str(c),
            )
            continue

        context = _text(context_row.get(c))
        if context not in CONTEXTS:
            issues.blocking(
                "INVALID_YEAR_BLOCK",
                f"Column {c} (year {year}) must hold one of {CONTEXTS} in row 2, "
                f"found {context_row.get(c)!r}",
                sheet=tree.worksheet, row=2, column=str(c),
            )
            continue

        block = blocks_by_year.get(year)
        if block is None:
            block = YearBlock(year=year, columns={})
            blocks_by_year[year] = block
            tree.year_blocks.append(block)

        if context in block.columns:
            issues.blocking(
                "DUPLICATE_YEAR_CONTEXT",
                f"Year {year} already has a {context!r} column at {block.columns[context]}; "
                f"column {c} duplicates it",
                sheet=tree.worksheet, row=2, column=str(c),
            )
            continue

        block.columns[context] = c


def _read_tree_sheet(ws, tree: ParsedTree, issues: IssueCollector) -> None:
    rows: dict[int, dict[int, object]] = {}
    for r, row in enumerate(ws.iter_rows(min_row=1), start=1):
        rows[r] = {i: cell.value for i, cell in enumerate(row, start=1)}

    _read_year_blocks(rows, tree, issues)

    r = NODE_FIRST_DATA_ROW
    while r in rows:
        cells = rows[r]
        code = _text(cells.get(1))
        if code is None:
            break
        node = ParsedNode(
            node_code=code,
            parent_node_code=_text(cells.get(2)),
            display_order=_int(cells.get(3)),
            kpi_name=_text(cells.get(4)),
            description=_text(cells.get(5)),
            unit=_text(cells.get(6)),
            node_type=_text(cells.get(7)),
            formula=_text(cells.get(8)),
            favorable_direction=_text(cells.get(9)),
            row=r,
        )
        for block in tree.year_blocks:
            for context, column in block.columns.items():
                node.values.append(ParsedValue(
                    year=block.year,
                    context=context,
                    value=_decimal(cells.get(column), tree.worksheet, r, column, issues),
                    row=r,
                    column=column,
                ))
        tree.nodes.append(node)
        r += 1
