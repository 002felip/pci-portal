"""
Endpoint smoke test  —  DEV UTILITY.

Runs the whole API in-process against the JSON backend and asserts the things
that actually broke in v1 or that the rebuild changed:

  * every endpoint answers 200,
  * periods are DISCOVERED (not configured) and ordered correctly,
  * an unknown snapshot is a 400, never a silent substitution,
  * the full CRUD round-trip works, including registry coercion,
  * creating an initiative under a NEW snapshot code brings that period into
    existence and makes it comparable  -- the new-quarter path end to end,
  * a registry addition propagates to the DB column, the UI contract, the
    Pydantic models and completeness scoring.

    python backend/smoke_test.py
"""

from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

os.environ.setdefault("DATA_BACKEND", "json")

DATA = ROOT / "backend" / "data"
SANDBOX = ROOT / "backend" / "_smoke_data"

ok = 0
fail = 0


def check(label: str, condition: bool, detail: str = "") -> None:
    global ok, fail
    if condition:
        ok += 1
        print(f"  PASS  {label}{('  ' + detail) if detail else ''}")
    else:
        fail += 1
        print(f"  FAIL  {label}{('  ' + detail) if detail else ''}")


def main() -> int:
    if not (DATA / "initiatives.json").exists():
        print("No seed found. Run `python backend/seed_demo.py` first "
              "(or `make demo-seed`).")
        return 2

    # Work on a copy: the CRUD tests write, and the seed should survive.
    if SANDBOX.exists():
        shutil.rmtree(SANDBOX)
    shutil.copytree(DATA, SANDBOX)
    os.environ["SEED_DIR"] = str(SANDBOX)

    from fastapi.testclient import TestClient

    from backend.app.main import app

    c = TestClient(app)

    print("\n1. Every endpoint answers")
    for url in ["/api/health", "/api/meta", "/api/meta/periods", "/api/meta/fields",
                "/api/initiatives?limit=5", "/api/analytics/compare",
                "/api/analytics/top?limit=10", "/api/analytics/monthly",
                "/api/analytics/quarters?stack_by=stage_code",
                "/api/analytics/stage-matrix",
                "/api/analytics/tree?levels=bu_code&levels=lever_code",
                "/api/analytics/completeness"]:
        r = c.get(url)
        check(url, r.status_code == 200, f"-> {r.status_code}")

    print("\n2. Periods are discovered, not configured")
    p = c.get("/api/meta/periods").json()
    codes = [s["id"] for s in p["snapshots"]]
    check("snapshots found", len(codes) >= 1, str(codes))
    check("years found", len(p["years"]) >= 1, str(p["years"]))
    check("default is latest vs previous",
          p["default_snapshot"] == codes[-1],
          f'{p["default_baseline"]} -> {p["default_snapshot"]}')
    check("ordering is chronological, not lexical",
          codes == sorted(codes, key=lambda x: p and
                          [s["sort_key"] for s in p["snapshots"] if s["id"] == x][0]),
          str(codes))

    print("\n3. An unknown period is rejected, never substituted")
    check("unknown snapshot -> 400",
          c.get("/api/analytics/compare?snapshot=Q99").status_code == 400)

    print("\n4. Monthly window follows the snapshot (year + cutoff)")
    for s in p["snapshots"]:
        r = c.get(f"/api/analytics/monthly?snapshot={s['id']}")
        j = r.json() if r.status_code == 200 else {}
        end = {"M": s["period_ordinal"], "Q": 3 * (s["period_ordinal"] or 0),
               "H": 6 * (s["period_ordinal"] or 0), "FY": 12}[s["period_type"]]
        tgt = j.get("financial_target_cum_musd", [])
        check(f"monthly {s['period_key']}",
              r.status_code == 200 and j["year"] == s["fiscal_year"]
              and j["cutoff_index"] == end - 1
              and all(a <= b + 1e-6 for a, b in zip(tgt, tgt[1:])))

    # Milestone blocks are YTD: carried across gaps, summed across rows.
    from backend.app.analytics import build_monthly
    ms = lambda scen, m, v: {"metric_code": "MILESTONES", "scenario_code": scen,
                             "period_month": f"2026-{m:02d}-01", "value": v}
    series = build_monthly([
        {"metrics": [ms("TARGET", 4, 3), ms("TARGET", 6, 5), ms("ACTUAL", 4, 2)]},
        {"metrics": [ms("TARGET", 2, 1), ms("ACTUAL", 5, 1)]},
    ], 2026, cutoff_month=6)
    check("milestones carried forward and summed",
          series.milestones_target_cum[:7] == [0, 1, 1, 4, 4, 6, 6]
          and series.milestones_actual_cum[:7] == [0, 0, 0, 2, 3, 3, 3]
          and not any(series.financial_target_cum_musd)
          and not any(series.kpi_sample),
          f"{series.milestones_target_cum} {series.milestones_actual_cum}")

    print("\n5. CRUD round-trip with registry coercion")
    snap = p["default_snapshot"]
    payload = {
        "snapshot_id": snap, "name": "Smoke-test kiln retrofit", "bu_code": "SSG",
        "source_initiative_id": "SMOKE-001", "site": "Plant 2",
        "stage_code": "IMPLEMENTING", "track_code": "ON", "aligned_lobp": True,
        "lobp_lever": "Cost", "lever_code": "COST", "opex_capex_code": "OPEX",
        "owner": "A. Souza", "sponsor": "B. Lima", "kpi_name": "Unit cost",
        "kpi_uom": "USD/t", "kpi_baseline": 100, "kpi_target_cy": 90,
        "actions_total": 10, "actions_done": 3, "value_target_cy": 4_200_000,
        "value_realized_cy": 0, "confidence_code": "HIGH",
        "planned_end": "2026-11-30",
    }
    r = c.post("/api/initiatives", json=payload)
    check("create -> 201", r.status_code == 201, str(r.status_code))
    row = r.json()
    check("planned_end derived to year/quarter",
          row.get("planned_end_year") == 2026 and row.get("planned_end_quarter") == 4,
          f'{row.get("planned_end_year")} Q{row.get("planned_end_quarter")}')

    iid = row["id"]
    r = c.patch(f"/api/initiatives/{iid}",
                json={"value_target_cy": 5_500_000, "track_code": "at risk"})
    check("patch -> 200", r.status_code == 200)
    check("vocab coerced ('at risk' -> AT_RISK)",
          r.json().get("track_code") == "AT_RISK", str(r.json().get("track_code")))
    check("edited_fields recorded (survives next Excel load)",
          bool(r.json().get("edited_fields", {}).get("value_target_cy")))

    found = c.get(f"/api/initiatives?snapshot={snap}&search=Smoke-test").json()
    check("appears in filtered list", found["total"] == 1, f'total={found["total"]}')

    check("delete -> 200", c.delete(f"/api/initiatives/{iid}").status_code == 200)

    print("\n6. A brand-new quarter comes into existence from data alone")
    new_id = "snp_smoke_q9"
    r = c.post("/api/initiatives", json={**payload, "snapshot_id": new_id,
                                        "source_initiative_id": "SMOKE-NEW"})
    check(f"create under {new_id} -> 201", r.status_code == 201, str(r.status_code))
    p2 = c.get("/api/meta/periods").json()
    check(f"{new_id} now in the picker",
          new_id in [s["id"] for s in p2["snapshots"]],
          str([s["id"] for s in p2["snapshots"]]))
    check("becomes comparable with no config change",
          c.get(f"/api/analytics/compare?snapshot={new_id}&baseline={snap}"
                ).status_code == 200)
    c.delete(f"/api/initiatives/{r.json()['id']}")

    print("\n7. A snapshot can be withdrawn and brought back")
    all_snaps = c.get("/api/snapshots").json()
    check("GET /snapshots lists every snapshot",
          len(all_snaps) >= len(p["snapshots"]),
          f'{len(all_snaps)} total / {len(p["snapshots"])} picker-visible')

    # The plan baseline is the one thing archiving refuses to touch: the
    # "plan" comparison mode would silently lose its anchor.
    c.post(f"/api/snapshots/{snap}/plan-baseline")
    r = c.post(f"/api/snapshots/{snap}/archive")
    check("archiving the plan baseline -> 409", r.status_code == 409, str(r.status_code))
    c.delete(f"/api/snapshots/{snap}/plan-baseline")

    r = c.post(f"/api/snapshots/{snap}/archive")
    check("archive -> 200", r.status_code == 200, str(r.status_code))
    check("archived_at stamped", bool(r.json().get("archived_at")))
    check("re-archiving -> 409",
          c.post(f"/api/snapshots/{snap}/archive").status_code == 409)

    p3 = c.get("/api/meta/periods").json()
    check("archived snapshot leaves every picker",
          snap not in [s["id"] for s in p3["snapshots"]])
    check("archived snapshot stays in the administrative list",
          snap in [s["id"] for s in c.get("/api/snapshots").json()])
    # Archived follows locked semantics: hidden from pickers, still resolvable
    # by id, so an existing deep link keeps working.
    check("archived snapshot is still addressable by id",
          c.get(f"/api/analytics/compare?snapshot={snap}").status_code == 200)
    check("editing an archived snapshot -> 409",
          c.patch(f"/api/snapshots/{snap}", json={"label": "nope"}).status_code == 409)

    r = c.delete(f"/api/snapshots/{snap}/archive")
    check("restore -> 200", r.status_code == 200, str(r.status_code))
    check("archived_at cleared", r.json().get("archived_at") is None)
    check("restored snapshot is back in the picker",
          snap in [s["id"] for s in c.get("/api/meta/periods").json()["snapshots"]])
    check("restoring a live snapshot -> 409",
          c.delete(f"/api/snapshots/{snap}/archive").status_code == 409)

    print("\n8. Snapshot metadata is editable, identity is not")
    before = next(s for s in c.get("/api/snapshots").json() if s["id"] == snap)
    r = c.patch(f"/api/snapshots/{snap}",
                json={"label": "Smoke label", "notes": "Loaded from the smoke test."})
    check("patch label + notes -> 200", r.status_code == 200, str(r.status_code))
    check("label round-trips", r.json()["label"] == "Smoke label", r.json()["label"])
    check("notes round-trips", r.json()["notes"] == "Loaded from the smoke test.")
    check("identity untouched by a metadata edit",
          r.json()["period_key"] == before["period_key"]
          and r.json()["sort_key"] == before["sort_key"]
          and r.json()["revision"] == before["revision"])
    check("a blank label -> 400",
          c.patch(f"/api/snapshots/{snap}", json={"label": "   "}).status_code == 400)
    # An unknown key is dropped by the whitelist rather than written through.
    c.patch(f"/api/snapshots/{snap}", json={"sort_key": 1, "state": "locked"})
    check("derived fields are not patchable",
          c.get("/api/snapshots").json()
          and next(s for s in c.get("/api/snapshots").json()
                   if s["id"] == snap)["sort_key"] == before["sort_key"])

    c.post(f"/api/snapshots/{snap}/plan-baseline")
    check("unpinning the plan baseline -> 409",
          c.patch(f"/api/snapshots/{snap}", json={"pinned": False}).status_code == 409)
    c.delete(f"/api/snapshots/{snap}/plan-baseline")
    c.patch(f"/api/snapshots/{snap}",
            json={"label": before["label"], "notes": None,
                  "pinned": before["pinned"]})

    # The end-to-end upload stamp has no automated cover -- there is no
    # fixture workbook in the repo -- so assert the plumbing either side of it:
    # the field is served, defaults to null, and survives a write/read cycle.
    every = c.get("/api/snapshots").json()
    check("uploaded_at is exposed on every snapshot",
          all("uploaded_at" in s_ for s_ in every))
    from backend.app.repositories import get_repository
    for _repo in get_repository():
        _snap = next(s_ for s_ in _repo.list_snapshots() if s_.id == snap)
        _repo.upsert_snapshot({**_snap.model_dump(),
                               "uploaded_at": "2026-05-04T10:00:00+00:00"})
        break
    check("uploaded_at round-trips through the store",
          next(s_ for s_ in c.get("/api/snapshots").json()
               if s_["id"] == snap)["uploaded_at"] == "2026-05-04T10:00:00+00:00")
    # And that the upload path actually assembles it, without a workbook.
    from backend.app.periods import resolve_period
    _new = resolve_period([], 2030, "Q", 1).identity
    check("a freshly resolved period carries the key",
          "uploaded_at" in _new and _new["uploaded_at"] is None)

    print("\n9. The registry is the single source of truth")
    from backend.app.models import Initiative
    from backend.app.registry import INITIATIVE_FIELDS, S_FIN, VOCAB
    from backend.app.schemas import InitiativePatch

    cols = {col.name for col in Initiative.__table__.columns}
    reg_cols = {f.column for f in INITIATIVE_FIELDS}
    check("every registry field is a DB column", reg_cols <= cols,
          f"{len(reg_cols)} fields / {len(cols)} columns")
    check("every editable field is patchable",
          {f.column for f in INITIATIVE_FIELDS if f.editable}
          <= set(InitiativePatch.model_fields))
    fields = c.get("/api/meta/fields").json()
    check("UI contract matches the registry",
          {f["column"] for f in fields["fields"]} == reg_cols)
    comp = c.get("/api/analytics/completeness").json()
    check("completeness scores registry fields",
          comp["fields_scored"] == len([f for f in INITIATIVE_FIELDS if f.required]) + 2,
          f'scored={comp["fields_scored"]}')

    rules = c.get("/api/analytics/completeness/rules")
    check("completeness rules list every stage", rules.status_code == 200
          and {s["code"] for s in rules.json()["stages"]}
          == {t[0] for t in VOCAB["stage"].terms})
    fin_labels = {f.label for f in INITIATIVE_FIELDS if f.section == S_FIN}
    evaluating = [d for d in comp["initiatives_detail"]
                  if d["stage_code"] == "EVALUATING"]
    check("Evaluating initiatives are not asked for financials by default",
          not any(fin_labels & set(d["missing_fields"]) for d in evaluating),
          f"{len(evaluating)} evaluating")
    custom = rules.json()["defaults"]["EVALUATING"] + ["value_target_cy"]
    c.put("/api/settings/completeness.required",
          json={"payload": {"by_stage": {"EVALUATING": custom}}})
    comp_custom = c.get("/api/analytics/completeness").json()
    check("a saved per-stage rule changes what is scored",
          comp_custom["total"] == comp["total"] + len(evaluating),
          f'{comp["total"]} -> {comp_custom["total"]}')
    c.delete("/api/settings/completeness.required")
    check("deleting the rule restores the defaults",
          c.get("/api/analytics/completeness").json()["pct"] == comp["pct"])

    print("\n10. Top-N movement matches by Initiative ID, not name")
    from backend.app.analytics import build_top

    def _r(sid, name, value):
        return {"bu_code": "ONT", "source_initiative_id": sid, "name": name,
                "value_target_cy": value}
    top = build_top(
        [_r("ONT-001", "Crusher uptime", 9e6), _r("ONT-002", "Crusher uptime", 8e6),
         _r("ONT-003", "Old name", 7e6)],
        [_r("ONT-001", "Crusher uptime", 9e6), _r("ONT-002", "Crusher uptime", 8e6),
         _r("ONT-003", "New name", 7e6)], "a", "b", limit=3).business_units[0]
    by_id = {m.source_initiative_id: m for m in top.movements}
    check("same name, different IDs tracked separately",
          by_id["ONT-001"].status == by_id["ONT-002"].status == "carried"
          and len(top.movements) == 3, str([m.status for m in top.movements]))
    check("renamed initiative is carried, flagged renamed",
          by_id["ONT-003"].status == "carried" and by_id["ONT-003"].renamed
          and by_id["ONT-003"].prev_name == "Old name")

    def _ra(sid, target, ytd, fy):
        return {"bu_code": "ONT", "source_initiative_id": sid, "name": sid,
                "value_target_cy": target, "value_actual_ytd": ytd, "value_actual_fy": fy}
    rows = [_ra("ONT-001", 9e6, 1e6, 6e6), _ra("ONT-002", 8e6, 3e6, 4e6),
            _ra("ONT-003", 7e6, 2e6, 5e6)]
    curated = {"ONT": {"a_items": [{"name": "curated"}]}}
    ytd = build_top(rows, rows, "a", "b", limit=3, basis="actual_ytd", overrides=curated)
    fy = build_top(rows, rows, "a", "b", limit=3, basis="actual_fy", overrides=curated)
    order = lambda t: [i.source_initiative_id for i in t.business_units[0].b_items]
    check("actual bases rank and value by actuals, ignoring curated overrides",
          order(ytd) == ["ONT-002", "ONT-003", "ONT-001"]
          and order(fy) == ["ONT-001", "ONT-003", "ONT-002"]
          and ytd.business_units[0].b_items[0].value_musd == 3.0
          and fy.basis == "actual_fy", f"{order(ytd)} / {order(fy)}")

    print("\n10b. Actuals YTD and actuals + forecast come from the monthly ACTUAL block")
    from backend.app.analytics import annotate_actuals

    def _pt(month, value, scen="ACTUAL", year=2026):
        return {"metric_code": "FINANCIAL_VALUE", "scenario_code": scen,
                "period_month": f"{year}-{month:02d}-01", "value": value}
    row = {"metrics": [_pt(1, 10), _pt(2, 20), _pt(5, 50), _pt(9, 90), _pt(11, 110),
                       _pt(12, 999, "TARGET"), _pt(12, 777, year=2025)]}
    annotate_actuals([row], 2026, 6)
    check("YTD is the period-end figure, carried across gaps",
          row["value_actual_ytd"] == 50, str(row["value_actual_ytd"]))
    check("actuals + forecast is the last figure through December",
          row["value_actual_fy"] == 110, str(row["value_actual_fy"]))
    empty = annotate_actuals([{"metrics": []}], 2026, 6)[0]
    check("no monthly actuals reads as zero",
          empty["value_actual_ytd"] == empty["value_actual_fy"] == 0)

    items = c.get("/api/initiatives?limit=3").json()["items"]
    check("the register carries the derived actuals",
          all("value_actual_ytd" in i and "value_actual_fy" in i for i in items)
          and all(i["metrics"] == [] for i in items))
    whole = c.get("/api/analytics/compare").json()
    bu = whole["rows"][0]["bu_code"]
    one = c.get(f"/api/analytics/compare?filter=bu_code:{bu}").json()
    check("the filter bar narrows both sides of the comparison",
          [r["bu_code"] for r in one["rows"]] == [bu]
          and one["totals"]["a"]["total"] == whole["rows"][0]["a"]["total"]
          and one["totals"]["b"]["total"] == whole["rows"][0]["b"]["total"])
    check("comparison totals carry actuals",
          whole["totals"]["b"]["actual_ytd_musd"] > 0
          and "prioritized" not in whole["totals"]["b"],
          str(whole["totals"]["b"]))
    check("top-N accepts a filter and an actual basis",
          c.get(f"/api/analytics/top?basis=actual_fy&filter=bu_code:{bu}").json()["basis"]
          == "actual_fy")

    print("\n11. Import template round-trips through the parser")
    from io import BytesIO

    from openpyxl import load_workbook

    from backend.app.etl import parse_workbook
    from backend.app.registry import METRIC_BLOCKS, VOCAB, resolve_vocab

    r = c.get("/api/uploads/template")
    check("GET /uploads/template -> xlsx", r.status_code == 200
          and "spreadsheetml" in r.headers.get("content-type", ""))
    tpl = r.content
    parsed = parse_workbook(BytesIO(tpl), snapshot_id="tpl", base_year=2026)
    missing = [i.message for i in parsed.issues if "not found" in i.message]
    check("every required column is recognised", not missing, "; ".join(missing))
    check("parser picks the register sheet", parsed.sheet_name == "Initiative Register")
    check("every dropdown label resolves",
          all(resolve_vocab(d, label)[1] for d in VOCAB for _c, label, *_ in VOCAB[d].terms))
    wb = load_workbook(BytesIO(tpl))
    check("Existing IDs sheet lists IDs in use",
          wb["Existing IDs"].max_row > 10, f'rows={wb["Existing IDs"].max_row}')

    def workbook(rows: list[dict]) -> bytes:
        """The template with `rows` filled in, keyed by header label. Month
        headers repeat per block, so the first (6A, KPI target) wins."""
        book = load_workbook(BytesIO(tpl))
        ws = book["Initiative Register"]
        labels: dict = {}
        for i in range(1, ws.max_column + 1):
            labels.setdefault(ws.cell(3, i).value, i)
        for n, row in enumerate(rows, start=4):
            for label, value in row.items():
                ws.cell(n, labels[label], value)
        out = BytesIO()
        book.save(out)
        return out.getvalue()

    hist = c.get("/api/initiatives?limit=1").json()["items"][0]
    base = {"BU": hist["bu_code"], "Site / Area": "Plant 1", "Initiative Stage": "Implementing",
            "On / Off Track": "On Track", "Current Year Value Target": 1_000_000}
    raw = workbook([{**base, "Initiative ID": "X-1", "Initiative Name": "Smoke one",
                     "Mar": 42}])
    filled = parse_workbook(BytesIO(raw), snapshot_id="tpl", base_year=2026)
    check("a filled-in template parses", len(filled.records) == 1
          and filled.records[0]["value_target_cy"] == 1_000_000
          and filled.records[0]["stage_code"] == "IMPLEMENTING",
          str([i.message for i in filled.issues if i.severity == "error"]))
    check("monthly blocks are located",
          [(m.metric_code, m.scenario_code, m.period_month.isoformat(), m.value)
           for m in filled.records[0].metrics]
          == [("KPI", "TARGET", "2026-03-01", 42.0)],
          str(filled.records[0].metrics))

    # Both milestone blocks and Notes share the "8." banner; the keyword must
    # route each month to the right scenario.
    book = load_workbook(BytesIO(raw))
    ws = book["Initiative Register"]
    banner = ""
    for i in range(1, ws.max_column + 1):
        banner = ws.cell(1, i).value or banner
        if ws.cell(3, i).value == "Apr" and "MILESTONES" in banner:
            ws.cell(4, i, 3 if "TARGETS" in banner else 2)
    ws.cell(4, next(i for i in range(1, ws.max_column + 1)
                    if ws.cell(3, i).value == "Actions Completed"), 2)
    out = BytesIO()
    book.save(out)
    ms = parse_workbook(BytesIO(out.getvalue()), snapshot_id="tpl", base_year=2026)
    check("milestone blocks are located",
          [(m.metric_code, m.scenario_code, m.period_month.isoformat(), m.value)
           for m in ms.records[0].metrics if m.metric_code == "MILESTONES"]
          == [("MILESTONES", "TARGET", "2026-04-01", 3.0),
              ("MILESTONES", "ACTUAL", "2026-04-01", 2.0)],
          str(ms.records[0].metrics))
    check("actions completed is located", ms.records[0]["actions_done"] == 2)

    print("\n12. Initiative IDs are checked against earlier periods")
    form = {"fiscal_year": "2031", "period_type": "Q", "period_ordinal": "1",
            "dry_run": "true"}

    def upload(rows, **extra):
        files = {"file": ("smoke.xlsx", workbook(rows),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        resp = c.post("/api/uploads/excel", data={**form, **extra}, files=files)
        return resp.json() if resp.status_code == 200 else {"status": resp.status_code}

    bu, sid, name = hist["bu_code"], hist["source_initiative_id"], hist["name"]
    j = upload([{**base, "Initiative ID": sid, "Initiative Name": "Totally unrelated thing"}])
    check("reused ID with an unrelated name blocks the load",
          j["status"] == "failed"
          and [x["status"] for x in j["id_checks"]] == ["conflict"], str(j.get("status")))
    j = upload([{**base, "Initiative ID": sid, "Initiative Name": "Totally unrelated thing"}],
               confirmed_ids=f'["{bu}|{sid}"]')
    check("confirming the conflict lets it through",
          j["status"] == "dry_run" and j["id_checks"][0]["confirmed"], str(j.get("status")))
    j = upload([{**base, "Initiative ID": sid, "Initiative Name": f"{name} (phase 2)"}])
    check("a slight rename is informational only",
          j["status"] == "dry_run" and j["id_checks"][0]["status"] == "renamed")
    j = upload([{**base, "Initiative Name": name},
                {**base, "Initiative Name": "Brand new smoke initiative"}])
    assigned = [x["source_initiative_id"] for x in j.get("id_checks", [])
                if x["status"] == "assigned"]
    check("blank ID reuses the earlier ID of the same name", assigned[:1] == [sid],
          str(assigned))
    check("otherwise a fresh {BU}-NNN is minted",
          len(assigned) == 2 and assigned[1].startswith(f"{bu}-") and assigned[1] != sid,
          str(assigned))

    shutil.rmtree(SANDBOX, ignore_errors=True)

    print(f"\n{'=' * 52}\n  {ok} passed, {fail} failed\n{'=' * 52}")
    return 1 if fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
