"""
Mass-import template, generated from the registry.

The layout is exactly what etl.parse_workbook reads, so the template cannot
drift from the parser: a field added to INITIATIVE_FIELDS or a block added to
METRIC_BLOCKS appears here with no edit.

    Initiative Register   banners / Required tags / headers, then data rows,
                          with dropdowns for every vocabulary and Yes/No field
    Existing IDs          next free ID per BU + every ID already in use,
                          so nobody reuses one by accident
    Instructions          ID rules
    Lists (hidden)        dropdown sources
"""

from __future__ import annotations

import re
from io import BytesIO
from typing import Any

from .id_check import Key, format_id, next_free_ids
from .registry import (
    INITIATIVE_FIELDS, METRIC_BLOCKS, VOCAB, Kind, label_of,
)

REGISTER_SHEET = "Initiative Register"
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
DATA_ROWS = 500                          # rows that carry the dropdowns

_FILL = {"banner": "007E7A", "required": "EDB111", "recommended": "D9D9D9",
         "header": "E8F1F0"}
_BLOCK_TITLE = {("KPI", "TARGET"): "KPI MONTHLY TARGET",
                ("KPI", "ACTUAL"): "KPI MONTHLY ACTUAL",
                ("FINANCIAL_VALUE", "TARGET"): "FINANCIAL VALUE MONTHLY TARGET (YTD, USD)",
                ("FINANCIAL_VALUE", "ACTUAL"): "FINANCIAL VALUE MONTHLY ACTUAL (YTD, USD)",
                ("MILESTONES", "TARGET"): "MONTHLY ACTIONS/MILESTONES TARGETS",
                ("MILESTONES", "ACTUAL"): "MONTHLY ACTIONS/MILESTONES ACTUALS + FORECAST"}

INSTRUCTIONS = (
    "How to fill in the Initiative Register",
    "",
    "• Do not rename, reorder or delete the three header rows; data starts on row 4.",
    "• Initiative ID identifies an initiative within its BU across every quarter.",
    "    – Keep the same ID every quarter, even when the initiative is renamed.",
    "    – Never reuse an ID that belonged to a different initiative. "
    "See 'Existing IDs' for the IDs already taken and the next free one per BU.",
    "    – Leave it blank and the tool assigns the next free ID ({BU}-NNN); "
    "copy that ID back into your workbook so later quarters match.",
    "• On upload, an ID whose name differs a lot from earlier quarters is held "
    "for confirmation before anything is loaded.",
    "• Money columns are in USD (not $M). Monthly financial and milestone values are year-to-date.",
    "• Monthly blocks: CURRENT YEAR is the fiscal year chosen at upload.",
)


def _section_banner(section: str) -> str:
    """'1 · Identification' -> '1. IDENTIFICATION' (the workbook convention)."""
    m = re.match(r"\s*(\w+)\s*·\s*(.*)", section)
    return f"{m.group(1)}. {m.group(2).upper()}" if m else section.upper()


def _block_banner(block) -> str:
    title = _BLOCK_TITLE.get((block.metric_code, block.scenario_code),
                             f"{block.metric_code} {block.scenario_code}")
    year = "NEXT YEAR" if block.year_offset else "CURRENT YEAR"
    return f"{block.banner}. {title} – {year}"


def build_template(history: dict[Key, dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    def fill(key: str) -> PatternFill:
        return PatternFill("solid", fgColor=_FILL[key])

    wb = Workbook()
    ws = wb.active
    ws.title = REGISTER_SHEET

    # ---- hidden dropdown sources --------------------------------------------
    lists = wb.create_sheet("Lists")
    list_ref: dict[str, str] = {}
    domains = sorted({f.vocab for f in INITIATIVE_FIELDS if f.kind is Kind.VOCAB and f.vocab})
    for c, domain in enumerate([*domains, "__bool"], start=1):
        values = ["Yes", "No"] if domain == "__bool" else [t[1] for t in VOCAB[domain].terms]
        lists.cell(1, c, domain)
        for r, v in enumerate(values, start=2):
            lists.cell(r, c, v)
        col = get_column_letter(c)
        list_ref[domain] = f"Lists!${col}$2:${col}${len(values) + 1}"
    lists.sheet_state = "hidden"

    # ---- register: scalar fields ------------------------------------------------
    col = 0
    last_section = None
    for spec in INITIATIVE_FIELDS:
        col += 1
        if spec.section != last_section:
            ws.cell(1, col, _section_banner(spec.section))
            last_section = spec.section
        ws.cell(2, col, "Required" if spec.required else "Recommended") \
            .fill = fill("required" if spec.required else "recommended")
        head = ws.cell(3, col, spec.label)
        notes = [spec.help] if spec.help else []
        if spec.kind is Kind.MONEY:
            notes.append("USD, not $M.")
        if spec.kind is Kind.DATE:
            notes.append("Date (dd/mm/yyyy).")
        if notes:
            head.comment = Comment(" ".join(notes), "Strategic Cockpit")
        ws.column_dimensions[get_column_letter(col)].width = \
            40 if spec.kind is Kind.LONGTEXT else 18

        source = list_ref.get(spec.vocab or "") if spec.kind is Kind.VOCAB else \
            list_ref["__bool"] if spec.kind is Kind.BOOL else None
        if source:
            dv = DataValidation(type="list", formula1=f"={source}", allow_blank=True)
            dv.error = f"Pick a {spec.label} from the list."
            ws.add_data_validation(dv)
            letter = get_column_letter(col)
            dv.add(f"{letter}4:{letter}{DATA_ROWS + 3}")

    # ---- register: monthly blocks -------------------------------------------
    for block in METRIC_BLOCKS:
        for i, month in enumerate(MONTHS):
            col += 1
            if i == 0:
                ws.cell(1, col, _block_banner(block))
            ws.cell(2, col, "Recommended").fill = fill("recommended")
            ws.cell(3, col, month)
            ws.column_dimensions[get_column_letter(col)].width = 11

    for c in range(1, col + 1):
        ws.cell(1, c).fill = fill("banner")
        ws.cell(1, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(3, c).fill = fill("header")
        ws.cell(3, c).font = Font(bold=True)
        ws.cell(3, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 32
    ws.freeze_panes = "C4"

    # ---- existing IDs ---------------------------------------------------------
    ids = wb.create_sheet("Existing IDs")
    ids.append(["Next free ID per BU"])
    ids["A1"].font = Font(bold=True)
    ids.append(["BU", "Next free ID"])
    nxt = next_free_ids(history)
    for code, *_ in VOCAB["bu"].terms:
        ids.append([code, format_id(code, nxt.get(code, 1))])
    ids.append([])
    header_row = ids.max_row + 1
    ids.append(["BU", "Initiative ID", "Initiative Name", "Last period", "Stage"])
    for c in range(1, 6):
        ids.cell(header_row, c).font = Font(bold=True)
        ids.cell(header_row, c).fill = fill("header")
    for (bu, sid), h in sorted(history.items()):
        row = h["row"]
        ids.append([bu, sid, row.get("name"), h["period_key"],
                    label_of("stage", row.get("stage_code"))])
    for letter, width in zip("ABCDE", (8, 16, 60, 12, 18)):
        ids.column_dimensions[letter].width = width
    ids.freeze_panes = ids.cell(header_row + 1, 1)
    ids.auto_filter.ref = f"A{header_row}:E{ids.max_row}"

    # ---- instructions ---------------------------------------------------------
    info = wb.create_sheet("Instructions")
    for line in INSTRUCTIONS:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 120

    wb.move_sheet("Instructions", offset=-(len(wb.sheetnames) - 1))
    wb.active = wb.sheetnames.index(REGISTER_SHEET)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
