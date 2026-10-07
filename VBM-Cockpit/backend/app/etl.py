"""
Excel -> canonical records -> storage. Vocabulary, parsing, validation and the
idempotent load, in one module because they are one pipeline.

REGISTRY-DRIVEN
---------------
There is no column list, no header map and no coercion ladder in this file. The
parser asks the registry two questions -- "which field owns this header?" and
"how do I coerce a value of this Kind?" -- and the answers live in registry.py.
A renamed spreadsheet column is a one-line edit there and this file never moves.

Sheet layout (0-indexed):
    row 0  section banners  ("1. IDENTIFICATION", "6A. KPI MONTHLY ...")
    row 1  requirement tags ("Required" / "Recommended" / "Ignore")
    row 2  field names
    row 3+ data

Columns are located BY NAME and monthly blocks BY BANNER, so an inserted column
upstream does not break the load.

IDEMPOTENCY
-----------
`load_batch` is UNIQUE (content_sha256, snapshot_id). Re-uploading the same
bytes returns the original batch and writes nothing (status `duplicate`);
`force=True` overrides for a deliberate replay. Initiatives upsert on
(snapshot, bu, source_id) so a row keeps its UUID -- and therefore its metric
history and any FK references -- across reloads.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
from dataclasses import asdict, dataclass, field
from typing import Any, BinaryIO, Iterable

from .periods import enrich_snapshot
from .registry import (
    INITIATIVE_FIELDS, METRIC_BLOCKS, VOCAB, FieldSpec, Kind, coerce,
    match_header, normalise, resolve_vocab,
)

SHEET_CANDIDATES = ("initiativeregister",)
HEADER_ROW = 2
DATA_START_ROW = 3
MONTH_TOKENS = ("jan", "feb", "mar", "apr", "may", "jun",
                "jul", "aug", "sep", "oct", "nov", "dec")


def sha256_bytes(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


# =================================================================== types ===
@dataclass
class MetricPoint:
    metric_code: str
    scenario_code: str
    period_month: dt.date
    value: float | None


@dataclass
class InitiativeRecord:
    values: dict[str, Any] = field(default_factory=dict)     # column -> value
    extra: dict[str, Any] = field(default_factory=dict)
    metrics: list[MetricPoint] = field(default_factory=list)
    row_number: int = 0

    def __getitem__(self, key: str) -> Any:
        return self.values.get(key)

    def get(self, key: str, default: Any = None) -> Any:
        return self.values.get(key, default)


@dataclass
class ParseIssue:
    severity: str
    row: int | None
    field: str
    message: str

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


@dataclass
class ParseResult:
    records: list[InitiativeRecord]
    issues: list[ParseIssue]
    snapshot_id: str
    sheet_name: str
    base_year: int

    @property
    def errors(self) -> list[ParseIssue]:
        return [i for i in self.issues if i.severity == "error"]


@dataclass
class LoadReport:
    batch_id: str
    status: str
    snapshot_id: str
    source_name: str
    content_sha256: str
    period_key: str | None = None
    revision: int = 1
    state: str | None = None
    stored_filename: str | None = None
    original_filename: str | None = None
    initiatives_inserted: int = 0
    initiatives_updated: int = 0
    initiatives_deleted: int = 0
    metric_rows: int = 0
    carried_edits: int = 0
    carried_edits_flagged: int = 0
    warnings: list[dict] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    message: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


# ================================================================== parsing ==
def _pick_sheet(wb) -> Any:
    for name in wb.sheetnames:
        if normalise(name) in SHEET_CANDIDATES:
            return wb[name]
    for name in wb.sheetnames:
        if "register" in normalise(name):
            return wb[name]
    return wb[wb.sheetnames[0]]


def _forward_fill(values: Iterable[Any]) -> list[str]:
    out, last = [], ""
    for v in values:
        if v not in (None, ""):
            last = str(v)
        out.append(last)
    return out


def _banner_code(banner: str) -> str | None:
    m = re.match(r"\s*(\d[A-D]?)\s*[.\-]", banner or "")
    return m.group(1).upper() if m else None


def _month_of(header: Any) -> int | None:
    h = normalise(header)
    for i, token in enumerate(MONTH_TOKENS):
        if h.startswith(token):
            return i + 1
    m = re.match(r"^m?(\d{1,2})$", h)
    if m and 1 <= int(m.group(1)) <= 12:
        return int(m.group(1))
    return None


def parse_workbook(stream: BinaryIO, snapshot_id: str,
                   base_year: int | None = None) -> ParseResult:
    """
    `base_year` anchors the monthly blocks. It is DISCOVERED from the workbook's
    year headers when present, and only then falls back to the current calendar
    year -- the old build hard-coded 2026/2027 in the block table, which meant
    January 2027 silently loaded a year of data into the wrong columns.
    """
    from openpyxl import load_workbook

    wb = load_workbook(stream, data_only=True, read_only=True)
    ws = _pick_sheet(wb)
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    if len(grid) <= DATA_START_ROW:
        return ParseResult([], [ParseIssue("error", None, "sheet",
                                           "Sheet has no data rows.")],
                           snapshot_id, ws.title, base_year or dt.date.today().year)

    banners = _forward_fill(grid[0])
    headers = [normalise(h) for h in grid[HEADER_ROW]]
    raw_headers = [str(h or "") for h in grid[HEADER_ROW]]

    year = base_year or _discover_base_year(banners, raw_headers)
    issues: list[ParseIssue] = []
    written = _year_in_text(banners, raw_headers)
    if base_year and written and written != base_year:
        # Same Jan..Dec columns every year, so the dropdown decides the dates;
        # a year written in the sheet that disagrees is most likely a mis-pick.
        issues.append(ParseIssue(
            "warning", None, "fiscal_year",
            f"Workbook text mentions {written} but the snapshot is FY{base_year}; "
            f"monthly figures were dated {base_year}. Re-upload with the right "
            f"fiscal year if this is wrong."))

    # ---- locate scalar columns, by asking the registry -------------------
    col_of: dict[str, int] = {}
    for idx, h in enumerate(headers):
        if not h:
            continue
        for spec in INITIATIVE_FIELDS:
            if spec.column in col_of:
                continue
            if match_header(spec, h):
                col_of[spec.column] = idx
                break

    for spec in INITIATIVE_FIELDS:
        if spec.required and spec.column not in col_of and spec.excel:
            issues.append(ParseIssue("warning", None, spec.column,
                                     f"Required column '{spec.label}' not found; "
                                     f"it will read as missing for every row."))

    # ---- locate monthly blocks, by banner --------------------------------
    monthly: list[tuple[int, str, str, dt.date]] = []
    blocks = {b.banner: b for b in METRIC_BLOCKS}
    for idx, banner in enumerate(banners):
        code = _banner_code(banner)
        block = blocks.get(code or "")
        if not block:
            continue
        month = _month_of(raw_headers[idx])
        if month:
            monthly.append((idx, block.metric_code, block.scenario_code,
                            dt.date(year + block.year_offset, month, 1)))

    # ---- rows -------------------------------------------------------------
    records: list[InitiativeRecord] = []
    seen: dict[tuple[str, str], int] = {}

    for r_idx in range(DATA_START_ROW, len(grid)):
        row = grid[r_idx]
        excel_row = r_idx + 1

        def cell(column: str) -> Any:
            i = col_of.get(column)
            return row[i] if i is not None and i < len(row) else None

        name = coerce(_spec("name"), cell("name"))
        bu_code, bu_ok = resolve_vocab("bu", cell("bu_code"))
        if not name and not bu_code:
            continue                      # genuinely blank row
        if not name:
            issues.append(ParseIssue("error", excel_row, "name",
                                     "Initiative Name is blank."))
            continue
        if not bu_code:
            issues.append(ParseIssue("error", excel_row, "bu_code",
                                     f"BU '{cell('bu_code')}' not recognised."))
            continue
        if not bu_ok:
            issues.append(ParseIssue("warning", excel_row, "bu_code",
                                     f"BU '{cell('bu_code')}' fell back to {bu_code}."))

        values: dict[str, Any] = {}
        for spec in INITIATIVE_FIELDS:
            raw = cell(spec.column)
            if spec.kind is Kind.VOCAB:
                code, ok = resolve_vocab(spec.vocab or "", raw)
                if raw not in (None, "") and not ok:
                    issues.append(ParseIssue(
                        "warning", excel_row, spec.column,
                        f"{spec.label}: '{raw}' is not in the {spec.vocab} "
                        f"vocabulary; recorded as {code}."))
                values[spec.column] = code
            else:
                values[spec.column] = coerce(spec, raw)

        sid = values.get("source_initiative_id") or f"{bu_code}-R{excel_row}"
        key = (bu_code, str(sid))
        if key in seen:
            # A single initiative carrying several KPIs. Suffix rather than drop.
            n = seen[key] + 1
            seen[key] = n
            sid = f"{sid}#{n}"
            issues.append(ParseIssue("warning", excel_row, "source_initiative_id",
                                     f"Duplicate ID for {bu_code}; stored as {sid}."))
        else:
            seen[key] = 1
        values["source_initiative_id"] = str(sid)

        if values.get("planned_end"):
            d = values["planned_end"]
            values["planned_end_year"] = d.year
            values["planned_end_quarter"] = (d.month - 1) // 3 + 1

        points = []
        for idx, metric, scenario, period in monthly:
            raw = row[idx] if idx < len(row) else None
            val = coerce(_MONEYISH, raw)
            if val is not None:
                points.append(MetricPoint(metric, scenario, period, val))

        extra = {raw_headers[i]: row[i] for i in range(len(raw_headers))
                 if raw_headers[i] and i not in set(col_of.values())
                 and i not in {m[0] for m in monthly}
                 and i < len(row) and row[i] not in (None, "")}

        rec = InitiativeRecord(values=values, extra=extra, metrics=points,
                               row_number=excel_row)
        _reconcile_phasing(rec, year, issues)
        records.append(rec)

    return ParseResult(records, issues, snapshot_id, ws.title, year)


_MONEYISH = FieldSpec("__v", "value", Kind.NUMBER, "x")


def _spec(name: str) -> FieldSpec:
    return next(f for f in INITIATIVE_FIELDS if f.name == name)


def _year_in_text(banners: list[str], headers: list[str]) -> int | None:
    """First 4-digit year in the banners/headers, skipping next-year blocks
    (6B/6D) so their year+1 label is not mistaken for the base year."""
    next_year = {b.banner for b in METRIC_BLOCKS if b.year_offset}
    texts = [b for b in banners if _banner_code(b) not in next_year]
    texts += [h for h, b in zip(headers, banners) if _banner_code(b) not in next_year]
    for text in texts:
        m = re.search(r"(20\d{2})", text or "")
        if m:
            return int(m.group(1))
    return None


def _discover_base_year(banners: list[str], headers: list[str]) -> int:
    """Find a 4-digit year in the section banners; fall back to today."""
    return _year_in_text(banners, headers) or dt.date.today().year


def _reconcile_phasing(rec: InitiativeRecord, year: int,
                       issues: list[ParseIssue]) -> None:
    """The monthly financial target is year-to-date, so its final month should
    equal the annual target. 2% tolerance."""
    annual = rec.get("value_target_cy")
    if not annual:
        return
    points = [m for m in rec.metrics
              if m.metric_code == "FINANCIAL_VALUE"
              and m.scenario_code == "TARGET" and m.period_month.year == year]
    if not points:
        return
    last = max(points, key=lambda m: m.period_month)
    ytd = last.value or 0
    if ytd and abs(ytd - annual) > max(abs(annual) * 0.02, 1000):
        issues.append(ParseIssue(
            "warning", rec.row_number, "value_target_cy",
            f"YTD phasing reaches {ytd:,.0f} by {last.period_month:%b} but the "
            f"annual target is {annual:,.0f} ({(ytd - annual) / annual:+.1%})."))


# =============================================================== serialise ===
def records_to_json(records: list[InitiativeRecord], snapshot_id: str,
                    batch_id: str) -> list[dict]:
    """Canonical dicts for the JSON backend -- same keys the SQL rows expose."""
    out = []
    for r in records:
        row = dict(r.values)
        row["snapshot_id"] = snapshot_id
        row["id"] = f"{snapshot_id}:{row.get('bu_code')}:{row['source_initiative_id']}"
        row["extra"] = r.extra or None
        row["load_batch_id"] = batch_id
        row["metrics"] = [{"metric_code": m.metric_code,
                           "scenario_code": m.scenario_code,
                           "period_month": m.period_month.isoformat(),
                           "value": m.value} for m in r.metrics]
        out.append(row)
    return out


# ============================================================ SQL load =======
def load_parse_result(session, parsed: ParseResult, raw_bytes: bytes,
                      source_name: str, snapshot_id: str,
                      dry_run: bool = False, force: bool = False,
                      snapshot_identity: dict | None = None) -> LoadReport:
    """
    `snapshot_identity` is the fully-resolved identity dict (see
    `periods.enrich_snapshot`) computed once by the caller; when omitted, an
    identity is derived from `snapshot_id` directly (the legacy path, ahead of
    the structured period builder).
    """
    from sqlalchemy import select

    from .models import (
        Initiative, InitiativeMetricMonthly, LoadBatch, Snapshot,
    )

    digest = sha256_bytes(raw_bytes)
    warnings = [i.as_dict() for i in parsed.issues if i.severity == "warning"]
    errors = [i.as_dict() for i in parsed.errors]
    identity = snapshot_identity or enrich_snapshot(snapshot_id, snapshot_id=snapshot_id)
    base = dict(status="", batch_id="", snapshot_id=snapshot_id,
                period_key=identity.get("period_key"), revision=identity.get("revision", 1),
                state=identity.get("state"), stored_filename=identity.get("stored_filename"),
                original_filename=identity.get("original_filename"),
                source_name=source_name, content_sha256=digest,
                warnings=warnings, errors=errors)

    if errors:
        return LoadReport(**{**base, "status": "failed",
                             "message": f"{len(errors)} blocking error(s); nothing written."})

    existing = session.scalar(select(LoadBatch).where(
        LoadBatch.content_sha256 == digest,
        LoadBatch.snapshot_id == snapshot_id))
    if existing and not force:
        return LoadReport(**{**base, "batch_id": str(existing.id), "status": "duplicate",
                             "message": "Identical bytes already loaded for this "
                                        "snapshot; nothing written."})
    if dry_run:
        return LoadReport(**{**base, "status": "dry_run",
                             "initiatives_inserted": len(parsed.records),
                             "metric_rows": sum(len(r.metrics) for r in parsed.records),
                             "message": "Validation only - no rows written."})

    seed_vocabulary(session)
    snapshot_row = session.get(Snapshot, snapshot_id)
    if snapshot_row is None:
        session.add(Snapshot(**{k: v for k, v in identity.items() if k != "as_of_date"}))
    else:
        for k, v in identity.items():
            if k not in ("id", "as_of_date"):
                setattr(snapshot_row, k, v)
    # The parent row goes in on its own flush. There is no relationship()
    # between LoadBatch and Snapshot, so the unit of work does not know one
    # INSERT depends on the other and orders the mappers by name -- load_batch
    # before snapshot, which the FK rejects on a brand-new period.
    session.flush()

    batch = LoadBatch(source_name=source_name, content_sha256=digest,
                      snapshot_id=snapshot_id, row_count=len(parsed.records),
                      status="running")
    session.add(batch)
    session.flush()

    current = {(r.bu_code, r.source_initiative_id): r for r in session.scalars(
        select(Initiative).where(Initiative.snapshot_id == snapshot_id)).all()}

    # Cross-snapshot curation carry-forward (spec step 5.1): the most recent
    # OTHER published/locked snapshot's edited row for a natural key with no
    # equivalent hand-edit in THIS snapshot yet.
    prior_rows = session.execute(
        select(Initiative, Snapshot.sort_key, Snapshot.state)
        .join(Snapshot, Snapshot.id == Initiative.snapshot_id)
        .where(Initiative.snapshot_id != snapshot_id,
              Initiative.edited_fields.isnot(None),
              Snapshot.state.in_(("published", "locked")))).all()
    carry_candidates: dict[tuple[str, str], tuple[int, Initiative]] = {}
    for prior_row, sort_key, _state in prior_rows:
        pkey = (prior_row.bu_code, prior_row.source_initiative_id)
        if pkey not in carry_candidates or sort_key > carry_candidates[pkey][0]:
            carry_candidates[pkey] = (sort_key, prior_row)

    inserted = updated = metric_rows = carried = carried_flagged = 0
    keep: set[tuple[str, str]] = set()

    for rec in parsed.records:
        key = (rec.get("bu_code"), rec.get("source_initiative_id"))
        keep.add(key)
        row = current.get(key)
        is_new = row is None
        if row is None:
            row = Initiative(snapshot_id=snapshot_id)
            session.add(row)
            inserted += 1
        else:
            updated += 1

        if is_new and key in carry_candidates:
            _, old = carry_candidates[key]
            new_edited: dict[str, Any] = {}
            for column, entry in (old.edited_fields or {}).items():
                if not entry:
                    continue
                baseline = entry.get("source_value") if isinstance(entry, dict) else None
                incoming = rec.get(column)
                unchanged = isinstance(entry, dict) and incoming == baseline
                new_edited[column] = {
                    "source_value": baseline if unchanged else incoming,
                    "needs_review": not unchanged,
                }
                setattr(row, column, getattr(old, column))     # the curated value wins
            if new_edited:
                row.edited_fields = new_edited
                carried += 1
                if any(v.get("needs_review") for v in new_edited.values()):
                    carried_flagged += 1

        # Hand-edits from the Initiatives tab are never clobbered by a reload.
        protected = set((row.edited_fields or {}).keys())
        for column, value in rec.values.items():
            if column in protected:
                continue
            setattr(row, column, value)
        for derived in ("planned_end_year", "planned_end_quarter"):
            if derived in rec.values and derived not in protected:
                setattr(row, derived, rec.values[derived])
        row.extra = rec.extra or None
        row.load_batch_id = batch.id
        session.flush()

        # The file is the system of record for phasing: rewrite per initiative.
        session.query(InitiativeMetricMonthly).filter_by(initiative_id=row.id).delete()
        for m in rec.metrics:
            session.add(InitiativeMetricMonthly(
                initiative_id=row.id, metric_code=m.metric_code,
                scenario_code=m.scenario_code, period_month=m.period_month,
                value=m.value, load_batch_id=batch.id))
            metric_rows += 1

    # A snapshot mirrors the file rather than accumulating ghosts.
    removed = 0
    for key, row in current.items():
        if key not in keep:
            session.delete(row)
            removed += 1

    batch.status = "loaded"
    batch.completed_at = dt.datetime.now(dt.timezone.utc)
    session.flush()

    return LoadReport(**{**base, "batch_id": str(batch.id), "status": "loaded",
                         "initiatives_inserted": inserted,
                         "initiatives_updated": updated,
                         "initiatives_deleted": removed, "metric_rows": metric_rows,
                         "carried_edits": carried, "carried_edits_flagged": carried_flagged,
                         "message": f"{inserted} inserted, {updated} updated, "
                                    f"{removed} removed, {metric_rows} monthly points."})


def seed_vocabulary(session) -> None:
    """Push registry vocabulary into the tables. Safe to re-run."""
    from .models import VocabAlias, VocabTerm

    for dom in VOCAB.values():
        for code, label, order, colour in dom.terms:
            row = session.get(VocabTerm, (dom.domain, code))
            if row is None:
                session.add(VocabTerm(domain=dom.domain, code=code, label=label,
                                      sort_order=order, color_hex=colour))
            else:
                row.label, row.sort_order, row.color_hex = label, order, colour
        for alias, code in dom.aliases.items():
            if alias and session.get(VocabAlias, (dom.domain, alias)) is None:
                session.add(VocabAlias(domain=dom.domain, alias_norm=alias, code=code))
    session.flush()
